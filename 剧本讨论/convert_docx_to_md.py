"""Convert 剧本 docx files to Markdown for AI reading."""
import docx
import re
import os

def detect_heading(text, is_bold, font_size):
    """Return heading level (0=not heading, 1=#, 2=##, 3=###) or 0 for body."""
    if not is_bold:
        return 0
    text = text.strip()
    if not text:
        return 0

    # Large font = major section header
    if font_size and font_size >= 140000:
        return 1

    # Week markers: W0, W1, W10 ... W20, 第N周
    if re.match(r'^W\d{1,2}$', text) or re.match(r'^W\d{1,2}[—\-\s（(]', text):
        return 2
    if re.match(r'^第\d+周', text):
        return 2

    # Scene labels like: bg xxx, BGMxxx, cg xxx
    if re.match(r'^(bg|BGM|CG|//)', text, re.IGNORECASE):
        return 0

    # Short bold lines that look like section titles
    if len(text) <= 30:
        # Not a character dialogue line
        if re.match(r'^[^\s：:]{1,6}[：:].+', text):
            return 0
        # Numbered reference lists like "1=平静，2=害羞" or "15=委屈"
        if re.match(r'^\d+=', text):
            return 0
        # Metadata lines like "预期触发条件：..."
        if re.match(r'^(预期|触发|条件|备注|注意)', text):
            return 0
        return 3  # sub-heading

    return 0


def convert_docx_to_md(input_path, output_path):
    doc = docx.Document(input_path)
    lines = []
    prev_empty = False

    for p in doc.paragraphs:
        text = p.text

        # Get formatting from first run
        is_bold = False
        font_size = 0
        for r in p.runs:
            if r.bold:
                is_bold = True
            if r.font.size:
                font_size = r.font.size
                break

        # Skip completely empty paragraphs (but track for spacing)
        if not text.strip():
            if not prev_empty:
                lines.append('')
                prev_empty = True
            continue
        prev_empty = False

        # Detect scene directions
        stripped = text.strip()

        # Check for comment lines
        if stripped.startswith('//'):
            lines.append(f'<!-- {stripped[2:].strip()} -->')
            continue

        # Check for bg/cg/BGM lines (BGM must come before bg to avoid "BGM6" → "BG"+"M6")
        bg_match = re.match(r'^(BGM\d*|bg|cg)[：:\s]*(.*)', stripped, re.IGNORECASE)
        if bg_match:
            tag = bg_match.group(1).upper()
            rest = bg_match.group(2).strip()
            if rest:
                lines.append(f'> **{tag}**: {rest}')
            else:
                lines.append(f'> **{tag}**')
            continue

        # Detect headings
        level = detect_heading(text, is_bold, font_size)
        if level == 1:
            lines.append(f'\n# {stripped}\n')
        elif level == 2:
            lines.append(f'\n## {stripped}\n')
        elif level == 3:
            lines.append(f'\n### {stripped}\n')
        else:
            lines.append(stripped)

    # Write output
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))

    print(f'Converted: {input_path} → {output_path}')
    print(f'  Paragraphs: {len(doc.paragraphs)}, Output lines: {len(lines)}')


if __name__ == '__main__':
    base = 'F:/game/galgame/WOTA/剧本讨论'

    # File 1: 剧本（完整设定）(1) — newer version
    src1 = os.path.join(base, '剧本（完整设定） (1).docx')
    dst1 = os.path.join(base, '剧本（完整设定）.md')
    if os.path.exists(src1):
        convert_docx_to_md(src1, dst1)
    else:
        print(f'Not found: {src1}')

    # File 2: 剧本（完整设定，时间线排序）
    src2 = os.path.join(base, '剧本（完整设定，时间线排序）.docx')
    dst2 = os.path.join(base, '剧本（完整设定，时间线排序）.md')
    if os.path.exists(src2):
        convert_docx_to_md(src2, dst2)
    else:
        print(f'Not found: {src2}')
